"""ЗАМЕРЫ_УСТОЙЧИВОСТИ.xlsx — таблица Excel, которую просит задание (29.09).

Задание: «строки — 24 показателя, столбцы — 5 прогонов; для каждого отчёта посчитайте, сколько показателей совпали
во всех пяти прогонах, совпал ли вывод, минимум и максимум суммы, какие показатели разошлись». Здесь на каждый
отчёт свой лист с оценками пяти прогонов, а совпадения, суммы и выводы посчитаны формулами Excel; лист «Сводка»
собирает их формулами же. Оценки прогонов — данные из raw/ (build_submission.py), больше ничего не вписано руками.

openpyxl записывает формулы без значений, и программы просмотра (Quick Look, превью в почте) показывают пустые
ячейки. Поэтому значение каждой формулы считается и здесь и вписывается в файл как сохранённый результат; при
открытии Excel пересчитывает всё сам (fullCalcOnLoad). Совпадение формул с этими значениями проверено LibreOffice.
Даты внутри файла фиксированы, чтобы пересборка совпадала со сданным файлом побайтно.
"""
import io
import re
import zipfile
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.hyperlink import Hyperlink

STAMP = (2026, 9, 29, 12, 0, 0)          # дата сдачи: ставится всем частям файла вместо времени сборки
FONT = "Arial"
SCORE = '+0;"−"0;0'
HEAD_FILL = PatternFill("solid", fgColor="EEF2F7")
FRAGILE_FILL = PatternFill("solid", fgColor="FDE8DF")
LINE = Border(bottom=Side(style="thin", color="C8CDD3"))
FIRST, N = 6, 24                         # первая строка показателей на листе отчёта и их число
LAST = FIRST + N - 1
SUM_ROW, VERDICT_ROW = LAST + 1, LAST + 2
STABLE_ROW, SAME_ROW, MIN_ROW, MAX_ROW = LAST + 4, LAST + 5, LAST + 6, LAST + 7
RUN_COLS = "CDEFG"


def verdict(s):
    return "strong" if s > 12 else "weak" if s < -12 else "mixed"


class Book:
    """Workbook + значения формул, посчитанные тем же правилом, что и формула."""

    def __init__(self):
        self.wb = Workbook()
        self.cache = {}                   # (лист, ячейка) → значение формулы

    def put(self, ws, ref, value, bold=False, fmt=None, wrap=False, size=None, center=False):
        c = ws[ref]
        c.value = value
        c.font = Font(name=FONT, bold=bold, size=size or 10)
        if fmt:
            c.number_format = fmt
        c.alignment = Alignment(wrap_text=wrap, vertical="top", horizontal="center" if center else None)
        return c

    def formula(self, ws, ref, formula, value, **kw):
        self.cache[(ws.title, ref)] = value
        return self.put(ws, ref, "=" + formula, **kw)


def report_sheet(book, ws, r, names):
    """Лист отчёта: 24 показателя × 5 прогонов, итог, совпадение; внизу суммы, выводы и замеры."""
    s = r["устойчивость"]
    put, formula = book.put, book.formula
    put(ws, "A1", f"{r['период']}: оценки пяти прогонов", bold=True, size=12)
    put(ws, "A2", f"Исходный файл: {r['файл']}")
    put(ws, "A3", "Прогоны отличаются только порядком показателей в запросе. Итог — медиана пяти оценок (совпадает с "
                  "самым частым значением). «нет» в колонке I — хрупкая клетка, её нужно проверить человеку.")
    heads = ["Блок", "Показатель", "Прогон 1", "Прогон 2", "Прогон 3", "Прогон 4", "Прогон 5", "Итог (медиана)",
             "Совпали во всех 5", "Кто ставит"]
    for col, h in zip("ABCDEFGHIJ", heads):
        c = put(ws, f"{col}5", h, bold=True, wrap=True)
        c.fill, c.border = HEAD_FILL, LINE
    for k, x in enumerate(r["показатели"]):
        row = FIRST + k
        put(ws, f"A{row}", x["блок"])
        put(ws, f"B{row}", names[x["id"]])
        for col, v in zip(RUN_COLS, x["голоса"]):
            put(ws, f"{col}{row}", v, fmt=SCORE)
        votes = sorted(x["голоса"])
        formula(ws, f"H{row}", f"MEDIAN(C{row}:G{row})", votes[2], bold=True, fmt=SCORE)
        formula(ws, f"I{row}", f'IF(MAX(C{row}:G{row})=MIN(C{row}:G{row}),"да","нет")',
                "да" if votes[0] == votes[-1] else "нет", center=True)
        put(ws, f"J{row}", x["кто_ставит"])
        assert votes[2] == x["оценка"]
    ws.conditional_formatting.add(f"A{FIRST}:J{LAST}", FormulaRule(formula=[f'$I{FIRST}="нет"'], fill=FRAGILE_FILL))
    sums = s["сумма_по_прогонам"]
    assert sums == [sum(x["голоса"][k] for x in r["показатели"]) for k in range(5)]
    put(ws, f"B{SUM_ROW}", "Сумма 24 оценок", bold=True)
    put(ws, f"B{VERDICT_ROW}", "Вывод (сумма > 12 — strong, < −12 — weak)", bold=True)
    for col, v in zip(RUN_COLS + "H", sums + [r["сумма"]]):
        c = formula(ws, f"{col}{SUM_ROW}", f"SUM({col}{FIRST}:{col}{LAST})", v, bold=True, fmt=SCORE)
        c.border = Border(top=Side(style="thin", color="C8CDD3"))
        formula(ws, f"{col}{VERDICT_ROW}", f'IF({col}{SUM_ROW}>12,"strong",IF({col}{SUM_ROW}<-12,"weak","mixed"))',
                verdict(v), center=True)
    stable = sum(not x["хрупкая"] for x in r["показатели"])
    assert stable == s["совпало_во_всех_5"] and (len({verdict(v) for v in sums}) == 1) == s["вывод_совпал_во_всех_5"]
    rows = [(STABLE_ROW, "Совпали во всех 5 прогонах (из 24)", f'COUNTIF(I{FIRST}:I{LAST},"да")', stable),
            (SAME_ROW, "Вывод одинаков во всех 5 прогонах", f'IF(COUNTIF(C{VERDICT_ROW}:G{VERDICT_ROW},'
             f'C{VERDICT_ROW})=5,"да","нет")', "да" if s["вывод_совпал_во_всех_5"] else "нет"),
            (MIN_ROW, "Сумма по прогонам: минимум", f"MIN(C{SUM_ROW}:G{SUM_ROW})", s["сумма_min"]),
            (MAX_ROW, "Сумма по прогонам: максимум", f"MAX(C{SUM_ROW}:G{SUM_ROW})", s["сумма_max"])]
    for row, label, f, v in rows:
        put(ws, f"B{row}", label, bold=True)
        formula(ws, f"C{row}", f, v, bold=True, fmt=None if isinstance(v, str) or row == STABLE_ROW else SCORE,
                center=isinstance(v, str))
    for col, w in zip("ABCDEFGHIJ", (16, 38, 9, 9, 9, 9, 9, 10, 10, 24)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = f"C{FIRST}"
    landscape(ws)


def summary_block(book, ws, top, rs, title, names):
    """Строки сводки по отчётам + итог; все числа — ссылки на листы отчётов."""
    put, formula = book.put, book.formula
    put(ws, f"A{top}", title, bold=True, size=11)
    heads = ["Отчёт", "Период", "Совпали во всех 5 (из 24)", "Доля", "Вывод одинаков во всех 5", "Итоговая сумма",
             "Вывод", "Сумма: минимум", "Сумма: максимум", "Разброс суммы", "Разошлись: показатель (5 прогонов)",
             "Контроль при T=0,7: совпали (из 24)"]
    for col, h in zip("ABCDEFGHIJKL", heads):
        c = put(ws, f"{col}{top + 1}", h, bold=True, wrap=True)
        c.fill, c.border = HEAD_FILL, LINE
    first = top + 2
    for k, r in enumerate(rs):
        row, s, q = first + k, r["устойчивость"], f"'{r['id']}'!"
        c = put(ws, f"A{row}", r["id"])
        c.hyperlink = Hyperlink(ref="", location=f"{q}A1")
        c.font = Font(name=FONT, size=10, color="1F5FAD", underline="single")
        put(ws, f"B{row}", r["период"])
        stable = s["совпало_во_всех_5"]
        formula(ws, f"C{row}", f"{q}C{STABLE_ROW}", stable)
        formula(ws, f"D{row}", f"C{row}/24", stable / 24, fmt="0%")
        formula(ws, f"E{row}", f"{q}C{SAME_ROW}", "да" if s["вывод_совпал_во_всех_5"] else "нет", center=True)
        formula(ws, f"F{row}", f"{q}H{SUM_ROW}", r["сумма"], fmt=SCORE)
        formula(ws, f"G{row}", f"{q}H{VERDICT_ROW}", r["вывод"], center=True)
        formula(ws, f"H{row}", f"{q}C{MIN_ROW}", s["сумма_min"], fmt=SCORE)
        formula(ws, f"I{row}", f"{q}C{MAX_ROW}", s["сумма_max"], fmt=SCORE)
        formula(ws, f"J{row}", f"I{row}-H{row}", s["сумма_max"] - s["сумма_min"])
        div = "; ".join(f"{names[x['id']]} ({' '.join(f'{v:+d}'.replace('-', '−') if v else '0' for v in x['голоса'])})"
                        for x in r["показатели"] if x["хрупкая"]) or "—"
        put(ws, f"K{row}", div, wrap=True)
        put(ws, f"L{row}", s["совпало_при_T07"])
    last, tot = first + len(rs) - 1, first + len(rs)
    n = len(rs)
    stable = sum(r["устойчивость"]["совпало_во_всех_5"] for r in rs)
    put(ws, f"B{tot}", f"Итого по {n} отчётам", bold=True)
    c = formula(ws, f"C{tot}", f"SUM(C{first}:C{last})", stable, bold=True)
    c.border = Border(top=Side(style="thin", color="C8CDD3"))
    formula(ws, f"D{tot}", f"C{tot}/({n}*24)", stable / (n * 24), bold=True, fmt="0.0%")
    formula(ws, f"E{tot}", f'COUNTIF(E{first}:E{last},"да")',
            sum(r["устойчивость"]["вывод_совпал_во_всех_5"] for r in rs), bold=True, center=True)
    put(ws, f"F{tot}", f"из {n} отчётов — вывод одинаков")
    formula(ws, f"L{tot}", f"SUM(L{first}:L{last})", sum(r["устойчивость"]["совпало_при_T07"] for r in rs), bold=True)
    return tot


def landscape(ws):
    """Печать: альбомная страница, вся ширина листа на одной странице."""
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def cached(xml, sheet_cache):
    """Вписать посчитанные значения в пустые <v/> формул листа."""
    def sub(m):
        ref = m.group(1)
        v = sheet_cache[ref]
        if isinstance(v, str):
            return f'<c r="{ref}"{m.group(2)} t="str"><f>{m.group(3)}</f><v>{escape(v)}</v></c>'
        return f'<c r="{ref}"{m.group(2)}><f>{m.group(3)}</f><v>{float(v):.15g}</v></c>'
    out = re.sub(r'<c r="([A-Z]+\d+)"([^>]*)><f>(.*?)</f><v\s*/></c>', sub, xml)
    assert "<v />" not in out and "<v/>" not in out
    return out


def write(path, sber, other, names):
    book = Book()
    ws = book.wb.active
    ws.title = "Сводка"
    book.put(ws, "A1", "Замеры устойчивости: 15 пресс-релизов Сбера, по 5 прогонов", bold=True, size=13)
    book.put(ws, "A2", "Модель Qwen 3.8 27B (4 бита, локально), спецификация v7, температура 0; в каждом прогоне свой "
                       "порядок показателей в запросе. Оценки прогонов — на листах отчётов (ссылка в колонке A); все "
                       "числа этого листа, кроме колонок K и L, — формулы по этим листам.")
    book.put(ws, "A3", "Колонка K перечисляет показатели, оценки которых разошлись; L — контрольный прогон того же "
                       "метода при температуре 0,7 (данные, не формула).")
    end = summary_block(book, ws, 5, sber, "Сбер", names)
    summary_block(book, ws, end + 2, other, "Другие банки: проверка переноса, в основную сдачу не входят", names)
    for col, w in zip("ABCDEFGHIJKL", (14, 22, 11, 7, 11, 10, 9, 10, 10, 9, 70, 12)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "C7"
    landscape(ws)
    for r in sber + other:
        report_sheet(book, book.wb.create_sheet(r["id"]), r, names)
    book.wb.properties.creator = "Андрей Гладышев"
    book.wb.properties.title = "Замеры устойчивости"
    buf = io.BytesIO()
    book.wb.save(buf)
    sheets = {ws.title: f"xl/worksheets/sheet{k + 1}.xml" for k, ws in enumerate(book.wb.worksheets)}
    per = {}
    for (title, ref), v in book.cache.items():
        per.setdefault(sheets[title], {})[ref] = v
    stamp = "%04d-%02d-%02dT%02d:%02d:%02dZ" % STAMP
    src = zipfile.ZipFile(buf)
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename in per:
                data = cached(data.decode(), per[info.filename]).encode()
            elif info.filename == "docProps/core.xml":
                data = re.sub(r"\d{4}-\d\d-\d\dT\d\d:\d\d:\d\dZ", stamp, data.decode()).encode()
            z.writestr(zipfile.ZipInfo(info.filename, date_time=STAMP), data, zipfile.ZIP_DEFLATED)
    path.write_bytes(out.getvalue())
